from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
MANUSCRIPT_MD = ROOT / "manuscript_aki_dynamic_prediction_deployment_gated.md"
OUT_DOCX = ROOT / "manuscript_aki_dynamic_prediction_deployment_gated_submission.docx"
PACKAGE_MANIFEST = ROOT / "submission_package_manifest.md"
COVER_LETTER = ROOT / "cover_letter_bmj_health_care_informatics.md"
TRIPOD_CHECKLIST = ROOT / "tripod_ai_readiness_checklist.md"
PROBAST_SUMMARY = ROOT / "probast_ai_risk_of_bias_summary.md"
OUTPUT_ROOT = ROOT / "aki_ssl_survival_output_clinical_deployment_gated_v4"
FIG_DIR = OUTPUT_ROOT / "figures" / "publication"
REPORT_DIR = OUTPUT_ROOT / "reports"


INK = RGBColor(31, 77, 120)
BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(30, 30, 30)
MUTED = RGBColor(90, 90, 90)
LIGHT_FILL = "F4F6F9"
TABLE_FILL = "F4F6F9"
TABLE_BORDER = "B8C2CC"
PAGE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(run, name: str = "Calibri", size: float | None = None,
                 color: RGBColor | None = None, bold: bool | None = None,
                 italic: bool | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.333

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, INK, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ["List Bullet", "List Number"]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208

    if "Reference" not in styles:
        ref = styles.add_style("Reference", 1)
    else:
        ref = styles["Reference"]
    ref.font.name = "Calibri"
    ref._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    ref._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    ref.font.size = Pt(9)
    ref.paragraph_format.left_indent = Inches(0.28)
    ref.paragraph_format.first_line_indent = Inches(-0.28)
    ref.paragraph_format.space_after = Pt(3)
    ref.paragraph_format.line_spacing = 1.05

    if "Caption" in styles:
        cap = styles["Caption"]
        cap.font.name = "Calibri"
        cap._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        cap._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        cap.font.size = Pt(9)
        cap.font.italic = True
        cap.font.color.rgb = MUTED
        cap.paragraph_format.space_before = Pt(4)
        cap.paragraph_format.space_after = Pt(8)


def set_section_geometry(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def set_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header_p = section.header.paragraphs[0]
    header_p.text = ""
    left = header_p.add_run("Deployment-gated AKI prediction manuscript")
    set_run_font(left, size=9, color=MUTED)
    footer_p = section.footer.paragraphs[0]
    footer_p.text = ""
    add_page_number(footer_p)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))

    tbl_grid = tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl.append(tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        tbl_grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"
        elem = borders.find(qn(tag))
        if elem is None:
            elem = OxmlElement(tag)
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "4")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), TABLE_BORDER)


def markdown_inline(text: str) -> str:
    text = text.replace("`", "")
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    return text.strip()


def add_body_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(markdown_inline(text))
    set_run_font(run, size=11, color=DARK)


def add_small_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(markdown_inline(text))
    set_run_font(run, size=9.5, color=MUTED, italic=True)


def compute_widths(headers: list[str]) -> list[int]:
    n = len(headers)
    if n == 1:
        return [PAGE_WIDTH_DXA]
    lower = [h.lower() for h in headers]
    weights = []
    for h in lower:
        if any(k in h for k in ["candidate", "subgroup", "level", "database", "baseline source", "race"]):
            weights.append(2.4)
        elif any(k in h for k in ["events", "landmarks", "stays", "alerts", "n"]):
            weights.append(1.0)
        elif any(k in h for k in ["auc", "auprc", "brier", "ece", "slope", "intercept", "threshold", "sensitivity", "specificity", "ppv", "npv"]):
            weights.append(0.95)
        else:
            weights.append(1.2)
    total = sum(weights)
    raw = [max(650, int(PAGE_WIDTH_DXA * w / total)) for w in weights]
    diff = PAGE_WIDTH_DXA - sum(raw)
    raw[-1] += diff
    return raw


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    headers = [markdown_inline(x) for x in rows[0]]
    data = [[markdown_inline(x) for x in r] for r in rows[1:]]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.allow_autofit = False
    widths = compute_widths(headers)
    set_table_geometry(table, widths)
    set_table_borders(table)
    header_cells = table.rows[0].cells
    set_repeat_header(table.rows[0])
    for i, h in enumerate(headers):
        set_cell_shading(header_cells[i], TABLE_FILL)
        p = header_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h)
        set_run_font(run, size=7.5 if len(headers) > 6 else 8.5, color=DARK, bold=True)
    for row in data:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            is_num = bool(re.fullmatch(r"[-+]?\d+(\.\d+)?", value.replace(",", "")))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_num or len(value) <= 8 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            set_run_font(run, size=7.25 if len(headers) > 6 else 8.5, color=DARK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_tables(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        line = lines[i].strip()
        cells = [c.strip() for c in line.strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def add_title_page(doc: Document, lines: list[str]) -> None:
    title = markdown_inline(lines[0].lstrip("#").strip())
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(title)
    set_run_font(run, size=20, color=INK, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(10)
    run = subtitle.add_run("Submission manuscript draft | Deployment-gated retrospective transportability study")
    set_run_font(run, size=12, color=MUTED, italic=True)

    meta_lines = []
    for line in lines[1:8]:
        if ":" in line:
            meta_lines.append(markdown_inline(line))
    for item in meta_lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        label, value = item.split(":", 1)
        r = p.add_run(label + ": ")
        set_run_font(r, size=10, color=DARK, bold=True)
        r = p.add_run(value.strip())
        set_run_font(r, size=10, color=DARK)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(8)
    note.paragraph_format.space_after = Pt(10)
    r = note.add_run(
        "Clinical-use status: blocked. The manuscript and deployment packet are suitable for research submission and hospital review, not bedside use."
    )
    set_run_font(r, size=10, color=RGBColor(155, 28, 28), bold=True)

    rule = doc.add_paragraph()
    p_pr = rule._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D7DBE2")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_figures(doc: Document) -> None:
    figures = [
        (
            "Figure 1. External AUROC by horizon.",
            FIG_DIR / "figure_1_auroc_by_horizon.png",
            "Line chart comparing external AUROC across prediction horizons for candidate AKI models.",
        ),
        (
            "Figure 2. Calibration summary across evaluated external candidates.",
            FIG_DIR / "figure_2_calibration_summary.png",
            "Calibration figure summarizing observed and predicted risk alignment across evaluated external AKI model candidates.",
        ),
        (
            "Figure 3. Threshold operating characteristics for the selected 72-hour external model.",
            FIG_DIR / "figure_3_threshold_operating_characteristics.png",
            "Threshold operating curve showing alert rate, sensitivity, and related operating characteristics for the selected 72-hour AKI model.",
        ),
    ]
    doc.add_heading("Figures", level=1)
    for caption, path, alt_text in figures:
        if not path.exists():
            add_small_note(doc, f"Missing figure file: {path}")
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        shape = run.add_picture(str(path), width=Inches(6.25))
        shape._inline.docPr.set("descr", alt_text)
        shape._inline.docPr.set("title", caption)
        cap = doc.add_paragraph(style="Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.add_run(caption)


def build_docx() -> None:
    raw = MANUSCRIPT_MD.read_text(encoding="utf-8")
    lines = raw.splitlines()
    doc = Document()
    set_section_geometry(doc)
    configure_styles(doc)
    set_header_footer(doc)
    add_title_page(doc, lines)

    i = 1
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|"):
            rows, i = parse_tables(lines, i)
            add_markdown_table(doc, rows)
            continue
        if stripped.startswith("# "):
            i += 1
            continue
        if stripped.startswith("## References"):
            add_figures(doc)
            doc.add_heading("References", level=1)
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(markdown_inline(stripped[3:]), level=1)
        elif stripped.startswith("### "):
            doc.add_heading(markdown_inline(stripped[4:]), level=2)
        elif re.match(r"^\d+\. ", stripped):
            p = doc.add_paragraph(style="Reference")
            run = p.add_run(markdown_inline(stripped))
            set_run_font(run, size=9, color=DARK)
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(markdown_inline(stripped[2:]))
            set_run_font(run, size=11, color=DARK)
        else:
            add_body_paragraph(doc, stripped)
        i += 1

    core = doc.core_properties
    core.title = "Deployment-gated dynamic prediction of severe acute kidney injury"
    core.subject = "AKI prediction, MIMIC-IV, eICU, clinical deployment gate"
    core.keywords = "AKI; machine learning; clinical prediction; transportability; calibration; deployment gate"
    core.author = "Terry Yu"
    core.comments = "Generated from verified Markdown manuscript and D-drive retrospective outputs."
    doc.save(OUT_DOCX)


def write_submission_docs() -> None:
    model_report = json.loads((REPORT_DIR / "model_selection_and_claims_report.json").read_text(encoding="utf-8"))
    gate = json.loads((REPORT_DIR / "clinical_deployment_gate_report.json").read_text(encoding="utf-8"))
    selected = model_report["selected_best_external_row"]
    refs = len(re.findall(r"^\d+\. ", MANUSCRIPT_MD.read_text(encoding="utf-8"), flags=re.M))

    def rel(path: Path) -> str:
        try:
            return str(path.relative_to(ROOT)).replace("\\", "/")
        except ValueError:
            return str(path).replace("\\", "/")

    COVER_LETTER.write_text(f"""# Cover Letter - BMJ Health & Care Informatics

Dear Editors,

Please consider the manuscript, "Deployment-Gated Dynamic Prediction of Severe Acute Kidney Injury in Critical Care: A Transportability and Clinical-Readiness Study," for BMJ Health & Care Informatics.

This paper is positioned as a health-informatics and clinical AI governance study, not as a ready bedside alert. The contribution is a rerunnable, deployment-gated AKI prediction workflow that explicitly tests common retrospective EHR failure modes: measurement-process bias, imperfect baseline creatinine, treatment/documentation confounding, repeated correlated landmarks, cross-database transportability, calibration, threshold burden, and non-causal interpretation of SHAP/PSM diagnostics.

The completed D-drive analysis used MIMIC-IV and eICU archives and found a best prespecified 72-hour external AUROC of {selected['auc']:.3f} (95% cluster-bootstrap CI {selected['auc_ci_low']:.3f}-{selected['auc_ci_high']:.3f}) for `{selected['prefix']}`. Because this did not meet the AUROC >=0.80 guardrail and no threshold met both sensitivity and alert-burden requirements, the deployment gate remained `{gate['status']}`. This negative-to-cautious result is central to the paper: it demonstrates how a publication workflow can prevent unsupported clinical deployment claims.

The manuscript includes {refs} references, a target-journal framing, calibration and threshold reporting, subgroup and baseline-source audits, first-landmark-per-stay sensitivity, and a draft hospital deployment evidence packet. The work should interest readers concerned with the safe translation of clinical AI from public ICU datasets into local health systems.

The authors confirm that any clinical use would require local prospective shadow validation, locked thresholds, governance approval, monitoring, incident response, regulatory review, and human-factors signoff.

Sincerely,

Terry Yu
""", encoding="utf-8")

    TRIPOD_CHECKLIST.write_text(f"""# TRIPOD+AI-Oriented Readiness Checklist

Status: draft self-audit for submission support. This is not an official completed EQUATOR form.

| Domain | Manuscript/package coverage | Status |
| --- | --- | --- |
| Title/abstract identify prediction model and AI/ML nature | Title and abstract describe dynamic prediction, AKI/RRT, ICU, MIMIC-IV/eICU, and deployment gate | covered |
| Background and intended use | Introduction and deployment-gate sections define retrospective research purpose and blocked clinical use | covered |
| Data source and eligibility | Methods and cohort table report MIMIC-IV/eICU sources, exclusions, baseline sCr, landmarks, and urine-output exclusions | covered |
| Outcome definition | Severe AKI/RRT and KDIGO-compatible criteria described | covered |
| Predictor definition and timing | Landmarking, actionability gap, feature families, and audit-only metadata described | covered |
| Missingness/measurement process | Urine-output coverage, measurement-process ablation, baseline-source audit, and documentation bias described | covered |
| Sample size and events | Cohort and external metric tables report stays, landmarks, and events | covered |
| Model development | CPU HGB/XGBoost models and optional SSL survival branch described | covered |
| Validation | Strict MIMIC-to-eICU and eICU-to-MIMIC external validation reported | covered |
| Performance measures | AUROC, AUPRC, Brier, ECE, calibration intercept/slope, thresholds, subgroup audit, and cluster bootstrap included | covered |
| Clinical utility | Threshold operating characteristics reported; no deployable threshold claimed | covered |
| Interpretability | SHAP/PSM bounded as exploratory and non-causal | covered |
| Deployment readiness | Deployment gate and draft hospital evidence packet included; gate remains blocked | covered |
| Reproducibility | Notebook, reports, figures, config, and environment manifest generated | covered |
""", encoding="utf-8")

    PROBAST_SUMMARY.write_text(f"""# PROBAST+AI-Oriented Risk-of-Bias Summary

This is a pragmatic self-audit to support peer review. It is not a substitute for independent PROBAST+AI assessment.

## Participants
Risk: moderate to high. Public ICU databases differ in time, geography, coding, hospital mix, and documentation. Adult ICU eligibility and exclusions were explicit, but prospective local representativeness is unproven.

Mitigation: strict cross-database validation, subgroup audit, baseline-source audit, and deployment gate requiring local prospective validation.

## Predictors
Risk: moderate to high. EHR measurements encode observation intensity, clinician concern, treatment decisions, and documentation workflows.

Mitigation: actionability gap, pre-landmark feature timing, treatment/documentation ablations, measurement-process ablations, and exclusion of race/ethnicity from model predictors.

## Outcome
Risk: moderate. Severe AKI/RRT is clinically meaningful but baseline creatinine and urine-output documentation are imperfect.

Mitigation: KDIGO-compatible phenotype, baseline-source indicators, urine-output observability requirement, early severe AKI/RRT exclusion, and baseline-source subgroup reporting.

## Analysis
Risk: moderate. Row-level dynamic landmarks produce correlated observations and may inflate apparent precision.

Mitigation: cluster-bootstrap confidence intervals, first-landmark-per-stay sensitivity, external validation, calibration reporting, and explicit claim guardrails.

## Overall Judgment
The retrospective research workflow is publication-ready as a transportability and governance study. It is not clinically deployable until local prospective shadow validation, threshold approval, monitoring, incident response, regulatory review, governance approval, and human-factors signoff are completed.
""", encoding="utf-8")

    PACKAGE_MANIFEST.write_text(f"""# AKI Publication and Deployment-Review Package

Generated: {date.today().isoformat()}

## Primary Submission Files
- Manuscript DOCX: `{rel(OUT_DOCX)}`
- Manuscript Markdown source: `{rel(MANUSCRIPT_MD)}`
- Cover letter draft: `{rel(COVER_LETTER)}`
- Target journal assessment: `{rel(ROOT / 'target_journal_assessment.md')}`
- TRIPOD+AI-oriented checklist: `{rel(TRIPOD_CHECKLIST)}`
- PROBAST+AI-oriented risk-of-bias summary: `{rel(PROBAST_SUMMARY)}`

## Rerunnable Analysis Files
- Notebook: `{rel(ROOT / 'aki_ssl_survival_sectioned_workflow_clinical_deployment_gated_v4_d_drive_publication.ipynb')}`
- Runner: `{rel(ROOT / 'run_aki_v4_notebook.py')}`
- Output directory: `{rel(OUTPUT_ROOT)}/`

## Key Reports
- Model selection and claims: `{rel(REPORT_DIR / 'model_selection_and_claims_report.md')}`
- Clinical deployment gate: `{rel(REPORT_DIR / 'clinical_deployment_gate_report.md')}`
- Clinical bias/transportability audit: `{rel(REPORT_DIR / 'clinical_bias_transportability_audit.md')}`
- Model card: `{rel(REPORT_DIR / 'clinical_model_card.md')}`

## Deployment Evidence Packet
- Folder: `{rel(REPORT_DIR / 'deployment_evidence')}/`
- Status: draft/not approved for clinical use
- Gate: `{gate['status']}`

## Primary Verified Result
- 72-hour selected external model: `{selected['prefix']}`
- AUROC: {selected['auc']:.3f} (95% CI {selected['auc_ci_low']:.3f}-{selected['auc_ci_high']:.3f})
- AUPRC: {selected['auprc']:.3f}
- Brier: {selected['brier']:.3f}
- ECE: {selected['ece']:.3f}
- Interpretation: no valid 72-hour external AUROC >=0.80 candidate; no clinical deployability claim.
""", encoding="utf-8")


if __name__ == "__main__":
    build_docx()
    write_submission_docs()
    print(OUT_DOCX)
